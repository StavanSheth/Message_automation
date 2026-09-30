"""Unit tests for LocalXlsxSource and BrowserSpreadsheetSource adapters."""

import pytest
import openpyxl
from backend.sources.xlsx.adapter import LocalXlsxSource, calculate_row_checksum
from backend.sources.browser_sheet.adapter import BrowserSpreadsheetSource
from backend.domain.enums import SourceAccessStatus, RepliedStatus
from backend.domain.errors import ValidationError, SourceAccessError


def test_xlsx_import_valid_sheet():
    adapter = LocalXlsxSource("tests/fixtures/sample_contacts.xlsx")
    assert adapter.validate_access() == SourceAccessStatus.ACCESSIBLE

    adapter.open()
    rows = adapter.read_records()
    assert len(rows) == 3

    # Row 1 verification
    r1 = rows[0]
    assert r1.name == "Alice Smith"
    assert r1.instagram_url == "https://instagram.com/alice_designer"
    assert r1.username == "alice_designer"
    assert r1.expected_followers == 1250
    assert r1.replied_status == RepliedStatus.UNKNOWN
    assert r1.followup_1_message is not None
    assert r1.followup_1_delay_seconds == 86400
    assert r1.followup_2_message is not None
    assert r1.followup_2_delay_seconds == 172800
    assert r1.checksum != ""

    # Row 2 verification
    r2 = rows[1]
    assert r2.name == "Bob Jones"
    assert r2.replied_status == RepliedStatus.NO
    assert r2.followup_2_message is None

    # Row 3 verification
    r3 = rows[2]
    assert r3.name == "Carol White"
    assert r3.replied_status == RepliedStatus.YES
    adapter.close()


def test_xlsx_missing_required_columns(tmp_path):
    invalid_file = str(tmp_path / "missing_cols.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Username"])  # Missing Instagram URL and Message
    ws.append(["John", "johnny"])
    wb.save(invalid_file)
    wb.close()

    adapter = LocalXlsxSource(invalid_file)
    with pytest.raises(ValidationError) as exc_info:
        adapter.read_records()
    assert "Missing required columns" in str(exc_info.value)


def test_xlsx_missing_required_cell_values(tmp_path):
    invalid_file = str(tmp_path / "missing_cell.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Instagram URL", "Message"])
    ws.append(["", "https://instagram.com/test", "Hello"])  # Empty name
    wb.save(invalid_file)
    wb.close()

    adapter = LocalXlsxSource(invalid_file)
    with pytest.raises(ValidationError) as exc_info:
        adapter.read_records()
    assert "'Name' is required" in str(exc_info.value)


def test_xlsx_update_record_writeback(tmp_path):
    test_file = str(tmp_path / "writeback.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Instagram URL", "Message", "Replied"])
    ws.append(["Dave", "https://instagram.com/dave", "Hi Dave", "UNKNOWN"])
    wb.save(test_file)
    wb.close()

    adapter = LocalXlsxSource(test_file)
    adapter.open()
    success = adapter.update_record(row_index=2, updates={"replied_status": "YES", "status": "SENT"})
    assert success is True
    adapter.close()

    # Re-open and verify changes in file
    wb2 = openpyxl.load_workbook(test_file)
    ws2 = wb2.active
    assert ws2.cell(row=2, column=4).value == "YES"
    # Added "Status" column
    assert ws2.cell(row=1, column=5).value == "Status"
    assert ws2.cell(row=2, column=5).value == "SENT"
    wb2.close()


def test_browser_spreadsheet_url_adapter():
    # Valid spreadsheet URL
    adapter_valid = BrowserSpreadsheetSource("https://docs.google.com/spreadsheets/d/abc123/edit")
    assert adapter_valid.validate_url() is True
    assert adapter_valid.validate_access() == SourceAccessStatus.ACCESSIBLE

    # Invalid URL syntax
    adapter_invalid = BrowserSpreadsheetSource("not-a-valid-url")
    assert adapter_invalid.validate_url() is False
    assert adapter_invalid.validate_access() == SourceAccessStatus.UNSUPPORTED_STRUCTURE

    # Custom access hook (e.g. login required)
    adapter_auth = BrowserSpreadsheetSource(
        "https://example.com/sheet",
        browser_hook=lambda url: SourceAccessStatus.LOGIN_REQUIRED,
    )
    assert adapter_auth.validate_access() == SourceAccessStatus.LOGIN_REQUIRED
